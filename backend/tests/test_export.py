"""Tests de la génération des exports Excel et CSV."""
from types import SimpleNamespace

from openpyxl import load_workbook
import io

from app.services.export import build_csv, build_xlsx


def make_doc(**overrides):
    data = dict(
        supplier={"name": "Fournisseur SARL", "registration_number": "FR123", "address": "1 rue Test"},
        invoice_number="F-001",
        invoice_date="2026-09-01",
        due_date="2026-10-01",
        currency="EUR",
        subtotal=100.0,
        tax=20.0,
        total=120.0,
        items=[{"description": "Prestation", "quantity": 1, "unit_price": 100.0, "tax_rate": 20, "total": 100.0}],
    )
    data.update(overrides.pop("extracted_data", {}))
    return SimpleNamespace(extracted_data=data, confidence=overrides.get("confidence", 0.9))


def test_build_xlsx_has_two_sheets_with_expected_rows():
    docs = [make_doc(), make_doc(extracted_data={"invoice_number": "F-002"})]
    content = build_xlsx(docs)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Factures", "Lignes"]

    invoices_sheet = wb["Factures"]
    assert invoices_sheet.max_row == 3  # en-tête + 2 factures
    assert invoices_sheet.cell(row=2, column=1).value == "Fournisseur SARL"

    items_sheet = wb["Lignes"]
    assert items_sheet.max_row == 3  # en-tête + 1 ligne par facture


def test_build_csv_contains_semicolon_delimited_header_and_bom():
    docs = [make_doc()]
    content = build_csv(docs)
    text = content.decode("utf-8")
    assert text.startswith("\ufeffFournisseur")
    assert "Fournisseur SARL" in text
    assert ";" in text.splitlines()[0]
