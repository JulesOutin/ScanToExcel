"""Génération des exports Excel (.xlsx, 2 feuilles) et CSV à partir d'un lot de documents."""
import csv
import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.models.document import Document

HEADER = [
    "Fournisseur", "N° TVA/SIRET", "Adresse", "N° facture", "Date facture",
    "Échéance", "Devise", "Sous-total", "TVA", "Total", "Confiance",
]
ITEM_HEADER = [
    "N° facture", "Fournisseur", "Description", "Quantité", "Prix unitaire", "Taux TVA %", "Total ligne",
]


def _invoice_row(doc: Document) -> list:
    d = doc.extracted_data or {}
    supplier = d.get("supplier", {})
    return [
        supplier.get("name", ""), supplier.get("registration_number", ""), supplier.get("address", ""),
        d.get("invoice_number", ""), d.get("invoice_date", ""), d.get("due_date", ""),
        d.get("currency", ""), d.get("subtotal", 0), d.get("tax", 0), d.get("total", 0),
        doc.confidence or d.get("confidence", 0),
    ]


def build_xlsx(documents: list[Document]) -> bytes:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Factures"
    ws1.append(HEADER)
    for doc in documents:
        ws1.append(_invoice_row(doc))
    for i, _ in enumerate(HEADER, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = 20

    ws2 = wb.create_sheet("Lignes")
    ws2.append(ITEM_HEADER)
    for doc in documents:
        d = doc.extracted_data or {}
        supplier_name = d.get("supplier", {}).get("name", "")
        for item in d.get("items", []):
            ws2.append([
                d.get("invoice_number", ""), supplier_name, item.get("description", ""),
                item.get("quantity", 0), item.get("unit_price", 0), item.get("tax_rate", 0),
                item.get("total", 0),
            ])
    for i, _ in enumerate(ITEM_HEADER, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(documents: list[Document]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(HEADER)
    for doc in documents:
        writer.writerow(_invoice_row(doc))
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM pour Excel FR
